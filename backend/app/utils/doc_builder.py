"""
Document builder — creates .docx research papers using python-docx.
Optionally converts to PDF via LibreOffice headless.
"""
import os
import re
import logging
import subprocess
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models import SectionDraft, SourceChunk

logger = logging.getLogger(__name__)


def build_docx(
    title: str,
    sections: dict[str, SectionDraft],
    chunks: list[SourceChunk],
    citation_style: str,
    outline_order: list[str],
    output_path: str,
) -> str:
    """Build a .docx research paper and save to *output_path*.

    Parameters
    ----------
    title : str
        Paper title (derived from topic).
    sections : dict
        Section name → SectionDraft mapping.
    chunks : list
        All source chunks (for building the reference list).
    citation_style : str
        "APA" or "IEEE".
    outline_order : list[str]
        Ordered list of section names (determines document order).
    output_path : str
        Full path to write the .docx file.

    Returns
    -------
    str
        The path to the generated .docx file.
    """
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    paragraph_format = style.paragraph_format
    paragraph_format.space_after = Pt(6)
    paragraph_format.line_spacing = 1.5

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    doc.add_paragraph()
    doc.add_page_break()

    for section_name in outline_order:
        section = sections.get(section_name)
        if not section:
            continue

        heading_level = 1
        if section_name.lower() in ("abstract",):
            heading_level = 1
        doc.add_heading(section_name, level=heading_level)

        content = section.content.strip()
        if content:
            paragraphs = content.split("\n\n")
            for para_text in paragraphs:
                para_text = para_text.strip()
                if para_text:
                    doc.add_paragraph(para_text)

    doc.add_page_break()
    doc.add_heading("References", level=1)

    cited_chunk_ids = set()
    for section in sections.values():
        cited_chunk_ids.update(section.citations_used)

    chunk_map = {c.id: c for c in chunks}
    references = []
    for i, chunk_id in enumerate(sorted(cited_chunk_ids), 1):
        chunk = chunk_map.get(chunk_id)
        if not chunk:
            continue
        if citation_style == "APA":
            ref = f"[{i}] {chunk.title or 'Untitled'}. Retrieved from {chunk.url}"
        else:
            ref = f"[{i}] \"{chunk.title or 'Untitled'},\" [Online]. Available: {chunk.url}"
        references.append(ref)

    for ref in references:
        doc.add_paragraph(ref)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    logger.info("DOCX saved to %s", output_path)
    return output_path


def convert_to_pdf(docx_path: str) -> str | None:
    """Convert .docx to .pdf using LibreOffice headless.

    Returns the PDF path on success, None if LibreOffice is not available.
    """
    try:
        output_dir = os.path.dirname(docx_path)
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", docx_path,
             "--outdir", output_dir],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode == 0:
            pdf_path = docx_path.rsplit(".", 1)[0] + ".pdf"
            if os.path.exists(pdf_path):
                logger.info("PDF saved to %s", pdf_path)
                return pdf_path
        logger.warning("LibreOffice conversion failed: %s", result.stderr)
        return None
    except FileNotFoundError:
        logger.info("LibreOffice not found — skipping PDF conversion")
        return None
    except Exception as exc:
        logger.warning("PDF conversion error: %s", exc)
        return None
